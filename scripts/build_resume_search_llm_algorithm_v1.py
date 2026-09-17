"""Build a search/recommendation and LLM-application algorithm resume.

The document preserves the V13 black-and-white one-page layout while changing
the evidence hierarchy for algorithm-oriented internships.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt


REFERENCE = Path(
    r"C:\Users\ming\Desktop\袁明珠_简历\袁明珠_Agent开发工程师Java_岗位定向简历_V13.docx"
)
OUTPUT = Path(
    r"C:\Users\ming\Desktop\袁明珠_简历\袁明珠_搜索推荐_LLM应用算法_岗位定向简历_V1.docx"
)


SEARCH_BULLETS = [
    (
        "数据集与基准：",
        "针对公开数据粒度与复杂购买意图脱节，构建三层评测：KuaiSearch 46,079 篇商品文档、507 条真实 Query、12 个品类；"
        "439 件二手手机、24 条复杂 Query；Amazon ESCI 50 条 Query、1,036 个官方标注 pair。",
    ),
    (
        "检索选型：",
        "对比 ES Standard/SmartCN、BGE Dense、Weighted RRF、Cross-Encoder 与 LLM 重排；"
        "混合方案 nDCG@10 +1.45pp 但 Recall@50 -5.01pp，Cross-Encoder 延迟约 4.9 倍，最终保留 ES Standard Top50。",
    ),
    (
        "约束理解与重排：",
        "将自然语言预算、品牌、成色与电池条件写入 TaskState，经 MySQL 权威核验、硬条件过滤和确定性重排；"
        "评测硬约束违规 0/831，ESCI 候选重排 nDCG@10=0.7069（随机基线 0.5283）。",
    ),
    (
        "商家评论 RAG：",
        "针对短评证据分散，设计结构化过滤、双路召回与内容重排，将离线 nDCG@5 从 0.368 提升至 0.788；"
        "因延迟与忠实度门未通过，仅保留 Shadow，未强行进入默认链路。",
    ),
    (
        "行为信号与记忆：",
        "区分显式偏好与低权威可过期点击信号，以 MySQL 版本链和 Redis 投影治理跨 Session 状态；"
        "425 用户未来行为离线评测 nDCG@10 +0.0074（95% CI [0.0037,0.0116]），全局默认仍关闭。",
    ),
    (
        "评测工程化：",
        "将 10 路候选池构建封装为显式 Skill 与本地 STDIO MCP，共用确定性内核、manifest/receipt/SHA 证据链；"
        "133 项 SHA 复核零失败，显式 Skill/MCP 输出字节一致。",
    ),
]


LLM_BULLETS = [
    (
        "Agent 架构选型：",
        "从固定 PAE 演进为 LangGraph durable 上的受约束 ReAct；24 场景每臂 33 轮，差异回答双评审盲评 ReAct/fixed/tie=7/0/3，保留 fixed_v1 回滚。",
    ),
    (
        "Context 与指代：",
        "构建类型化 ContextCompiler 和服务端 ReferenceContext；25 组长历史总 Token -22.9%，"
        "8 个真实会话 21 轮答案/证据 21/21 字节一致，13 个续问两次独立 AI 盲评均为平局。",
    ),
    (
        "Multi-Agent 协作：",
        "仅在证据缺口触发只读 EvidenceResearchAgent，以任务版本和候选域绑定消息，Redis 原子合并后由父 Agent 确定性渲染；"
        "公开开发集 5/9→9/9、冻结合成留出 8/8、原始子观测泄漏 0。",
    ),
    (
        "持久化执行：",
        "以 Redis Checkpoint、ToolInbox、租约和 fencing token 支持澄清挂起、主动暂停与崩溃续跑；"
        "130/130 恢复、12/12 篡改拒绝，进程内 RTO P95 约 75ms。",
    ),
    (
        "在线交易边界：",
        "Agent 仅生成候选域绑定提案，Java 回查 JWT 和确认后写入；Redis Lua、Idempotency-Key 与权威状态对账闭合重试，"
        "200 个任务争 50 库存时 50 成功、150 售罄、0 超卖。",
    ),
]


SKILLS = [
    (
        "检索与评测：",
        "Elasticsearch、BM25、Dense Retrieval、RRF、Cross-Encoder；Recall/nDCG/MRR/Hit@K、Bootstrap CI、消融实验与 qrels 候选池。",
    ),
    (
        "LLM / Agent：",
        "LangGraph、ReAct、RAG、Tool Calling、Context、Memory、Multi-Agent、Checkpoint、Skill/MCP。",
    ),
    (
        "工程能力：",
        "Python、Java、FastAPI、Spring Boot、MySQL、Redis、Kafka、Docker Compose、pytest、JUnit。",
    ),
]


def _replace(paragraph, label: str, text: str) -> None:
    if len(paragraph.runs) < 2:
        raise RuntimeError(f"paragraph lacks body run: {paragraph.text}")
    paragraph.runs[0].text = label
    paragraph.runs[1].text = text
    for run in paragraph.runs[2:]:
        run.text = ""


def _replace_whole(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def _delete(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def build() -> None:
    doc = Document(REFERENCE)
    doc.core_properties.title = "袁明珠 - 搜索推荐 / LLM 应用算法实习简历"
    doc.core_properties.subject = "搜索推荐与 LLM 应用算法岗位定向一页中文简历"

    search_slots = []
    llm_slots = []
    skill_slots = []

    search_labels = {
        "Agent 运行时：",
        "持久化执行与故障恢复：",
        "上下文与指代：",
        "检索与数据评测：",
        "商家评论 RAG：",
        "评测工程化：",
        "Multi-Agent：",
        "长期记忆：",
    }
    llm_labels = {
        "交易闭环：",
        "库存与订单补偿：",
        "异步投影：",
        "缓存与限流：",
        "架构取舍：",
    }

    for paragraph in list(doc.paragraphs):
        text = paragraph.text
        if text.startswith("Agent 开发工程师（Java）实习生"):
            suffix = text.split("  |  北京", 1)[1]
            _replace_whole(paragraph, "搜索推荐 / LLM 应用算法实习生  |  北京" + suffix)
        elif text.startswith("电商购物 Agent 与 Java 交易平台"):
            if len(paragraph.runs) >= 2:
                paragraph.runs[0].text = "电商搜索推荐与 LLM 导购 Agent  |  个人项目\t"
                for run in paragraph.runs[1:-1]:
                    run.text = ""
            else:
                _replace_whole(paragraph, "电商搜索推荐与 LLM 导购 Agent  |  个人项目\t2026.07 - 至今")
        elif text.startswith("面向多品类电商检索"):
            _replace_whole(
                paragraph,
                "面向多品类电商搜索与二手手机复杂约束导购，构建检索排序、商家评论 RAG、多轮决策和受控交易链路；"
                "以离线指标、配对实验及故障注入驱动方案选型。",
            )
        elif text == "智能购物 Agent":
            _replace_whole(paragraph, "搜索推荐与评测")
        elif text.startswith("技术栈：Python / FastAPI"):
            _replace_whole(
                paragraph,
                "技术栈：Python / Elasticsearch / BM25 / BGE / RRF / Cross-Encoder / MySQL / Redis",
            )
        elif paragraph.runs and paragraph.runs[0].text in search_labels:
            search_slots.append(paragraph)
        elif text == "Java 电商交易后端":
            _replace_whole(paragraph, "LLM 应用与在线工程")
        elif text.startswith("技术栈：Java 17"):
            _replace_whole(
                paragraph,
                "技术栈：Python / FastAPI / LangGraph / ReAct / RAG / MCP / Java / Spring Boot",
            )
        elif paragraph.runs and paragraph.runs[0].text in llm_labels:
            llm_slots.append(paragraph)
        elif paragraph.runs and paragraph.runs[0].text in {
            "Java / Spring：",
            "数据与中间件：",
            "Agent 工程：",
        }:
            skill_slots.append(paragraph)

    if len(search_slots) != 8 or len(llm_slots) != 5 or len(skill_slots) != 3:
        raise RuntimeError(
            f"slot mismatch: search={len(search_slots)}, llm={len(llm_slots)}, skills={len(skill_slots)}"
        )

    for paragraph, (label, text) in zip(search_slots[:6], SEARCH_BULLETS):
        _replace(paragraph, label, text)
    for paragraph in search_slots[6:]:
        _delete(paragraph)

    for paragraph, (label, text) in zip(llm_slots, LLM_BULLETS):
        _replace(paragraph, label, text)
    for paragraph, (label, text) in zip(skill_slots, SKILLS):
        _replace(paragraph, label, text)

    project_labels = {label for label, _ in SEARCH_BULLETS + LLM_BULLETS}
    for paragraph in doc.paragraphs:
        if paragraph.runs and paragraph.runs[0].text in project_labels:
            paragraph.paragraph_format.space_after = Pt(0.9)
            paragraph.paragraph_format.line_spacing = 1.02
            for run in paragraph.runs:
                run.font.size = Pt(9.5)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
