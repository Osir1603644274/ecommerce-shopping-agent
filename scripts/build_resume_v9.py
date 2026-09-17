"""Build the black-and-white, evidence-bound one-page Chinese resume V9."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUTPUT = Path(r"C:\Users\ming\Desktop\袁明珠_简历\袁明珠_Agent开发工程师Java_岗位定向简历_V9.docx")
BODY_CN = "宋体"
HEAD_CN = "黑体"
LATIN = "Times New Roman"


def set_run_font(run, size: float, *, bold: bool = False, east_asia: str = BODY_CN, latin: str = LATIN) -> None:
    run.bold = bold
    run.font.name = latin
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)
    color = run._element.get_or_add_rPr().find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        run._element.get_or_add_rPr().append(color)
    color.set(qn("w:val"), "000000")


def set_keep(paragraph, *, next_paragraph: bool = False) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:keepLines")) is None:
        ppr.append(OxmlElement("w:keepLines"))
    if next_paragraph and ppr.find(qn("w:keepNext")) is None:
        ppr.append(OxmlElement("w:keepNext"))


def add_bottom_border(paragraph, size: int = 8, space: int = 1) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    ppr.append(borders)


def add_bullet_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_id = max(
        [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))],
        default=0,
    ) + 1
    num_id = max([int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))], default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", "•"), ("w:lvlJc", "left")):
        node = OxmlElement(tag)
        attr = "w:val"
        node.set(qn(attr), value)
        level.append(node)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "250")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "250")
    indent.set(qn("w:hanging"), "170")
    ppr.append(tabs)
    ppr.append(indent)
    level.append(ppr)
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    rpr.append(rfonts)
    level.append(rpr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def apply_bullet(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    node = OxmlElement("w:numId")
    node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(node)
    ppr.append(num_pr)


def add_section(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2.8)
    p.paragraph_format.space_after = Pt(2.0)
    p.paragraph_format.line_spacing = 1.0
    set_keep(p, next_paragraph=True)
    add_bottom_border(p)
    set_run_font(p.add_run(text), 10.6, bold=True, east_asia=HEAD_CN, latin="Arial")


def add_subhead(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2.0)
    p.paragraph_format.space_after = Pt(0.7)
    p.paragraph_format.line_spacing = 1.0
    set_keep(p, next_paragraph=True)
    set_run_font(p.add_run(text), 9.35, bold=True, east_asia=HEAD_CN, latin="Arial")


def add_left_right(doc: Document, left: str, right: str, *, size: float = 9.3, after: float = 0.4) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.tab_stops.add_tab_stop(Cm(18.35), WD_TAB_ALIGNMENT.RIGHT)
    set_keep(p)
    set_run_font(p.add_run(left), size, bold=True, east_asia=HEAD_CN)
    p.add_run("\t")
    set_run_font(p.add_run(right), size)


def add_label_line(doc: Document, label: str, text: str, *, size: float = 8.9, after: float = 0.4) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    set_keep(p)
    set_run_font(p.add_run(label), size, bold=True, east_asia=HEAD_CN, latin="Arial")
    set_run_font(p.add_run(text), size, latin="Arial")


def add_bullet(doc: Document, num_id: int, label: str, text: str, *, size: float = 9.05, after: float = 0.55) -> None:
    p = doc.add_paragraph()
    apply_bullet(p, num_id)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.widow_control = True
    set_keep(p)
    set_run_font(p.add_run(label), size, bold=True, east_asia=HEAD_CN)
    set_run_font(p.add_run(text), size)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(0.72)
    section.bottom_margin = Cm(0.68)
    section.left_margin = Cm(1.02)
    section.right_margin = Cm(1.02)
    section.header_distance = Cm(0.2)
    section.footer_distance = Cm(0.2)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(9.05)
    normal._element.rPr.rFonts.set(qn("w:ascii"), LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CN)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    doc.core_properties.title = "袁明珠 - Agent 开发工程师（Java）实习简历 V9"
    doc.core_properties.subject = "岗位定向一页中文简历（问题-方案-数据-取舍）"
    doc.core_properties.author = "袁明珠"
    doc.core_properties.keywords = "Java, Agent, ReAct, LangGraph, Spring Boot, Redis, MySQL, RAG, Multi-Agent"

    bullet_num = add_bullet_numbering(doc)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(0.5)
    name.paragraph_format.line_spacing = 1.0
    set_run_font(name.add_run("袁明珠"), 18.5, bold=True, east_asia=HEAD_CN, latin="Arial")

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(0.8)
    contact.paragraph_format.line_spacing = 1.0
    set_run_font(
        contact.add_run("Agent 开发工程师（Java）实习  |  北京  |  17559557087  |  17559557087@163.com"),
        9.25,
        east_asia=BODY_CN,
        latin="Arial",
    )

    add_section(doc, "教育经历")
    add_left_right(doc, "北京航空航天大学  |  电子信息  硕士", "2025.09 - 2028.06", size=9.45)
    add_left_right(doc, "河北工业大学  |  电子信息工程  本科", "2021.09 - 2025.06", size=9.45)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0.2)
    p.paragraph_format.line_spacing = 1.0
    set_run_font(p.add_run("本科专业排名 1/126，推免至北京航空航天大学"), 9.1)

    add_section(doc, "项目经历")
    add_left_right(doc, "电商购物 Agent  |  个人项目", "2026.07 - 至今", size=9.55, after=0.45)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(0.35)
    intro.paragraph_format.line_spacing = 1.0
    set_run_font(
        intro.add_run("面向多轮商品检索、证据核验与受控交易：Agent 负责搜推决策，Java 后端持有身份、库存、订单与支付写边界。"),
        8.95,
    )

    add_subhead(doc, "智能购物 Agent")
    add_label_line(doc, "技术栈：", "Python / FastAPI / LangGraph / Redis / MySQL / Elasticsearch / MCP")
    add_bullet(
        doc,
        bullet_num,
        "运行时选型：",
        "针对固定 PAE 步骤多、推进僵硬，将其收敛为受约束 ReAct：模型选择签名动作，Executor/Validator 基于 TaskState 与 CandidateScope 执行校验；24 场景安全门 24/24，盲审 ReAct/fixed/tie=7/0/3，任务推进 3.3→4.9、实用性 3.0→5.0，保留 fixed_v1 回滚。",
    )
    add_bullet(
        doc,
        bullet_num,
        "崩溃恢复：",
        "针对“工具成功、下一 Checkpoint 未确认”导致重放的问题，以 Redis Checkpoint＋ToolInbox logical slot/lease/fence 约束执行；130/130 恢复、12/12 篡改拒绝，重复/缺失副作用与旧 Worker 覆盖均为 0，进程内 RTO P50/P95 42.182/75.475 ms。",
    )
    add_bullet(
        doc,
        bullet_num,
        "短期 Context：",
        "针对全量历史冗余及跨轮指代错绑，构建 ContextPack/typed ContextCompiler 与服务端 ReferenceContext 收据，绑定 session/task/revision/展示顺序/focus；25 组长历史两臂 25/25 等价，Prompt/总 Token 分别下降 26.43%/22.92%，真实网页验证越界与跨任务引用失败关闭。",
    )
    add_bullet(
        doc,
        bullet_num,
        "长期记忆：",
        "针对偏好跨会话持久化与可撤销治理，设计异步候选→用户确认→MySQL 权威版本链→Redis 投影，支持跨 Session 召回、纠正、停用与忘掉；真实三 Session 链路 42/42。",
    )
    add_bullet(
        doc,
        bullet_num,
        "商品检索选型：",
        "在 439 商品、24 Query 上对比 ES Standard/SmartCN、Dense、RRF、Cross-Encoder 与 LLM 重排；混合方案 nDCG@10 +1.45pp 但 Recall@50 -5.01pp，Cross-Encoder 延迟 4.9×、LLM 重排触发 6s 超时，最终保留 ES Standard Top50＋MySQL 核验＋确定性重排，硬约束违规 0/831。",
    )
    add_bullet(
        doc,
        bullet_num,
        "商家评论 RAG：",
        "针对词法/向量单路召回不足，采用 Vector/BM25 各 Top30、2:8 融合与内容重排，独立 test 的 nDCG@5 0.368→0.788；Context Selection 将评论 56.5→7.3 条、引用校验 100%，但重排平均 9.05s、完整回答忠实率 50%，因此保持 Shadow。",
    )
    add_bullet(
        doc,
        bullet_num,
        "Multi-Agent：",
        "针对单 Agent 证据调查不足，构建 Shopping Coordinator→只读 EvidenceResearchAgent→确定性合并，以 task/revision/CandidateScope 哈希和 Redis 原子 claim 通信；公开开发集 5/9→9/9、冻结合成留出 8/8、P95 5.50s，子观测泄漏 0。",
    )
    add_bullet(
        doc,
        bullet_num,
        "Skill / MCP：",
        "将人工候选池流程沉淀为 10 路检索确定性核心，并由显式 Skill 与本地 STDIO MCP 复用；记录 135 个 overlap、7 个独有候选及 17 个排序分歧，Core/Skill/MCP 输出字节一致。",
    )
    add_bullet(
        doc,
        bullet_num,
        "外部检索评测：",
        "在 Amazon ESCI 的 50 个 US 查询、1,036 对官方人工标注 query-product pairs 上验证候选重排，macro nDCG@10 0.7069（随机顺序 0.5283），配对差值 0.1786，95% CI [0.1163, 0.2430]。",
    )

    add_subhead(doc, "Java 电商交易后端")
    add_label_line(
        doc,
        "技术栈：",
        "Java 17 / Spring Boot / Spring Security / MyBatis / Spring Cloud / MySQL / Redis / Caffeine / Kafka / Elasticsearch",
    )
    add_bullet(
        doc,
        bullet_num,
        "模块与权威边界：",
        "按商品、库存、订单、支付/退款与搜索划分 Spring Boot 模块化单体；Agent 不持有订单写工具，下单需 Java 回查 JWT、校验候选范围并二次确认，GETDEL＋Idempotency-Key 分别防止确认重复消费与重复建单。",
    )
    add_bullet(
        doc,
        bullet_num,
        "库存与秒杀：",
        "Redis Lua 原子预扣/一人一单，MySQL 条件扣减＋唯一索引提供最终边界，Redis Stream 承接命令；严格 200 并发争 50 库存得到 50 成功、150 售罄、0 超卖。",
    )
    add_bullet(
        doc,
        bullet_num,
        "订单补偿：",
        "将 PENDING_PAYMENT→EXPIRED 条件迁移、库存/优惠券释放和过期 Outbox 纳入同一事务；2 个并发 Worker 仅 1 次状态迁移、1 条事件，库存 3→5、reserved 2→0；Stream 支持 Pending 接管、重试、DLQ 与幂等补偿。",
    )
    add_bullet(
        doc,
        bullet_num,
        "异步投影：",
        "订单/商品变更与 Outbox 同事务，Kafka→Inbox→Elasticsearch 以 projection cursor/external_gte 处理重复、乱序和旧版本事件，异常路径显式回到 MySQL 权威。",
    )
    add_bullet(
        doc,
        bullet_num,
        "缓存治理：",
        "商品详情采用 Caffeine＋Redis 两级缓存；事务提交后删除 Redis 并通过 Pub/Sub 广播跨实例 L1 失效，真实双实例验证旧缓存失效后读取 entityVersion=2。",
    )
    add_bullet(
        doc,
        bullet_num,
        "多维限流：",
        "以 Redis ZSET＋Lua 实现 IP/JWT 用户/归一化接口滑动窗口，避免固定窗口边界突刺和资源 ID 绕过；100 并发额度 20 时精确放行 20、拒绝 80。",
    )
    add_bullet(
        doc,
        bullet_num,
        "架构取舍：",
        "拆分 Gateway、Catalog/Search、Trade，引入 Feign、Bulkhead(8) 与 CircuitBreaker；正式场景 10/10，但同源对照拆分 P50/P95 增加 104.65%/11.88%、启动增加 158.64%，故保留模块化单体默认。",
    )
    add_bullet(
        doc,
        bullet_num,
        "工程验收：",
        "Java 后端 194/194、Gateway 4/4；Agent 全量回归 3593 passed、12 skipped、0 failed，覆盖并发、故障注入、幂等与恢复链路。",
    )

    add_section(doc, "专业技能")
    add_bullet(doc, bullet_num, "Java / Spring：", "熟练使用 Java 集合、异常、并发与 JVM 基础；掌握 Spring Boot、IOC/AOP、事务、Spring Security、MyBatis 与 Spring Cloud。", size=8.9, after=0.25)
    add_bullet(doc, bullet_num, "数据与中间件：", "掌握 MySQL 事务/隔离级别/索引/锁与 Redis 缓存、Lua、ZSET、Stream；具备 Kafka、Elasticsearch、Caffeine 项目实践。", size=8.9, after=0.25)
    add_bullet(doc, bullet_num, "Agent 与工程：", "掌握 ReAct/LangGraph、Tool Calling、Checkpoint、Context、RAG、Multi-Agent、长期记忆、Skill/MCP；使用 Maven、Docker Compose、JUnit、pytest 完成故障注入与证据化评测。", size=8.9, after=0)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
