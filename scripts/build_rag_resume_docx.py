from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = Path("outputs/resume/袁明珠-RAG检索算法实习-阿里千问.docx")

FONT_CN = "Microsoft YaHei"
FONT_LATIN = "Arial"
BLUE = RGBColor(24, 90, 173)
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(82, 92, 105)


def set_run_font(run, size=10.0, bold=False, color=INK, italic=False):
    run.font.name = FONT_LATIN
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), FONT_LATIN)
    fonts.set(qn("w:hAnsi"), FONT_LATIN)
    fonts.set(qn("w:eastAsia"), FONT_CN)


def set_style_font(style, size, bold=False, color=INK):
    style.font.name = FONT_LATIN
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), FONT_LATIN)
    fonts.set(qn("w:hAnsi"), FONT_LATIN)
    fonts.set(qn("w:eastAsia"), FONT_CN)


def add_style(doc, name, size, bold=False, color=INK, before=0, after=0, line=1.0):
    styles = doc.styles
    style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = styles["Normal"]
    set_style_font(style, size, bold=bold, color=color)
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    return style


def keep_with_next(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    keep = ppr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        ppr.append(keep)


def keep_together(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    keep = ppr.find(qn("w:keepLines"))
    if keep is None:
        keep = OxmlElement("w:keepLines")
        ppr.append(keep)


def set_cell_free_section_spacing(paragraph, before=0, after=0, line=1.0):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_section_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.add_run(text)
    keep_with_next(p)
    return p


def add_entry_line(doc, left_bold, left_rest, date_text):
    p = doc.add_paragraph(style="ResumeEntry")
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
    set_run_font(p.add_run(left_bold), size=9.8, bold=True)
    set_run_font(p.add_run(left_rest), size=9.3, color=MUTED)
    set_run_font(p.add_run("\t" + date_text), size=9.1, color=MUTED)
    keep_together(p)
    return p


def add_label_paragraph(doc, label, text, style="ResumeLine", label_color=INK):
    p = doc.add_paragraph(style=style)
    set_run_font(p.add_run(label), size=9.8, bold=True, color=label_color)
    set_run_font(p.add_run(text), size=9.8, color=INK)
    keep_together(p)
    return p


def configure_document():
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.5)

    normal = doc.styles["Normal"]
    set_style_font(normal, 10.0, color=INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.12

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, 12.0, bold=True, color=BLUE)
    h1.paragraph_format.space_before = Pt(7.0)
    h1.paragraph_format.space_after = Pt(3.2)
    h1.paragraph_format.line_spacing = 1.0
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, 10.0, bold=True, color=INK)
    h2.paragraph_format.space_before = Pt(2)
    h2.paragraph_format.space_after = Pt(1)
    h2.paragraph_format.line_spacing = 1.0

    add_style(doc, "ResumeName", 24, bold=True, color=INK, before=0, after=1.5, line=1.0)
    add_style(doc, "ResumeTarget", 11.0, bold=True, color=BLUE, before=0, after=2.5, line=1.0)
    add_style(doc, "ResumeMeta", 9.3, color=MUTED, before=0, after=1.2, line=1.0)
    add_style(doc, "ResumeEntry", 9.8, color=INK, before=0.5, after=1.2, line=1.0)
    add_style(doc, "ResumeLine", 9.8, color=INK, before=0, after=1.8, line=1.10)
    add_style(doc, "ResumeProject", 9.8, color=INK, before=0, after=3.2, line=1.16)

    return doc


def build_resume():
    doc = configure_document()

    p = doc.add_paragraph(style="ResumeName")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("袁明珠"), size=24, bold=True, color=INK)

    p = doc.add_paragraph(style="ResumeTarget")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("RAG 检索算法实习生  |  2028 届"), size=11.0, bold=True, color=BLUE)

    p = doc.add_paragraph(style="ResumeMeta")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        p.add_run("17559557087  |  1603644274@qq.com  |  北京"),
        size=9.3,
        color=MUTED,
    )

    p = doc.add_paragraph(style="ResumeMeta")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("每周至少到岗 4 天  |  可连续实习 6 个月"), size=9.3, color=MUTED)

    add_section_heading(doc, "教育经历")
    add_entry_line(doc, "北京航空航天大学", "  |  电子信息 · 硕士在读", "2025.09—2028.06")
    add_entry_line(doc, "河北工业大学", "  |  电子信息工程 · 本科", "2021.09—2025.06")

    add_section_heading(doc, "专业技能")
    add_label_paragraph(doc, "编程与工程：", "Python、Java、SQL；FastAPI、Spring Boot、MyBatis、Pydantic、Docker。")
    add_label_paragraph(doc, "大模型与 Agent：", "LangGraph、Function Calling、Prompt/Tool Schema、多轮 TaskState、Planner–Executor–Validator–Replanner。")
    add_label_paragraph(doc, "检索与评测：", "Elasticsearch/BM25、Qdrant、Dense Retrieval、RRF、Cross-Encoder；Hit@K、Recall@K、MRR、nDCG、P95。")
    add_label_paragraph(doc, "模型与数据：", "PyTorch、Transformers、BGE Embedding/Reranker；MySQL、Redis、Kafka。")

    add_section_heading(doc, "项目经历")
    add_entry_line(
        doc,
        "二手 3C Agentic Product Search 与自动化评测系统",
        "  |  个人项目",
        "2026.06—至今",
    )

    add_label_paragraph(
        doc,
        "项目概述：",
        "围绕二手 3C 多轮导购，在受控商品域内完成自然语言需求理解、检索召回、排序、权威事实核验与自动化评测闭环。",
        style="ResumeProject",
        label_color=BLUE,
    )

    add_label_paragraph(
        doc,
        "Agent 与 Query 理解：",
        "基于 FastAPI、LangGraph 和 Function Calling 构建有界 Agent，将自然语言需求沉淀为带 revision 的 TaskState；由 Planner、Executor、Validator、Replanner 完成检索、详情核验、比较、澄清与恢复闭环。",
        style="ResumeProject",
        label_color=BLUE,
    )
    add_label_paragraph(
        doc,
        "检索与排序：",
        "搭建并对比 Elasticsearch/BM25、BGE Dense Retrieval、RRF 融合和 Cross-Encoder 重排；将品牌、型号、内存、成色等硬约束与相关性排序分离，并回查 Java/MySQL 权威事实，防止陈旧索引突破业务约束。",
        style="ResumeProject",
        label_color=BLUE,
    )
    add_label_paragraph(
        doc,
        "RAG 与效果分析：",
        "使用 Qdrant、BGE Embedding 与 BM25 实现带来源引用的评论 RAG，拆分评估检索命中与回答忠实性；真实 Yelp 评论 25 题评测 Hit@3 为 60.0%，据失败样本确定 metadata filter、Hybrid Search、Query Rewrite 与 Rerank 优化方向。",
        style="ResumeProject",
        label_color=BLUE,
    )
    add_label_paragraph(
        doc,
        "数据与评测闭环：",
        "构建候选池、标注/qrels、冻结测试与自动化 Benchmark/Harness，统一输出 Hit@K、Recall、MRR、nDCG、硬约束违规和端到端延迟；覆盖多轮澄清、无结果、约束冲突、依赖故障和状态恢复，并生成可追溯 trace、评分报告与哈希绑定产物。",
        style="ResumeProject",
        label_color=BLUE,
    )

    doc.core_properties.title = "袁明珠 - RAG 检索算法实习简历"
    doc.core_properties.subject = "阿里千问日常算法实习生投递版"
    doc.core_properties.author = "袁明珠"
    doc.core_properties.keywords = "RAG, 检索算法, Agent, Elasticsearch, BM25, Qdrant"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build_resume()
