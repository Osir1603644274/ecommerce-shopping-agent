"""Build the final black-and-white one-page Chinese resume V13.

V13 keeps V12's complete Agent/Java split and updates only claims that gained
new bounded evidence on 2026-09-03: real multi-turn Context fidelity,
transaction-command reconciliation, and governed memory evaluation.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

import build_resume_v12


OUTPUT = Path(
    r"C:\Users\ming\Desktop\袁明珠_简历\袁明珠_Agent开发工程师Java_岗位定向简历_V13.docx"
)


BULLETS = {
    "持久化执行与故障恢复：": (
        "基于 Redis Checkpoint、ToolInbox、租约和 fencing token 支持澄清挂起、主动暂停与崩溃续跑；"
        "130/130 恢复、12/12 篡改拒绝，重复/缺失副作用与旧 Worker 覆盖均为 0，进程内 RTO P95 约 75ms。"
    ),
    "上下文与指代：": (
        "针对全量历史膨胀与跨轮错指，构建类型化 ContextCompiler 和服务端 ReferenceContext；"
        "25 组长历史总 Token -22.9%，8 个真实会话 21 轮答案/证据 21/21 字节一致，13 个续问双盲评均为平局；默认仍关闭。"
    ),
    "Multi-Agent：": (
        "仅在证据缺口触发只读 EvidenceResearchAgent，消息绑定任务版本与候选域，Redis 原子合并后由父 Agent 确定性渲染；"
        "公开开发集 5/9→9/9、冻结合成留出 8/8、原始子观测泄漏 0，异常回退单 Agent。"
    ),
    "长期记忆：": (
        "设计“异步候选→用户确认→MySQL 权威版本链→Redis 投影”，支持跨 Session 召回、纠正、停用与删除；"
        "真实三 Session 链路 42/42，425 用户未来行为上 nDCG@10 +0.0074（95% CI [0.0037,0.0116]）；默认关闭。"
    ),
    "交易闭环：": (
        "Agent 仅生成绑定 CandidateScope 的确认提案，Java 回查 JWT 后持有写权限；GETDEL 防重复确认，Idempotency-Key 防重复建单。"
    ),
    "库存与订单补偿：": (
        "Redis Lua 原子预扣配合 MySQL 条件扣减、用户/请求唯一索引与 Outbox/Kafka/Inbox；"
        "副作用完成但回执未写时回查 Java 权威状态并按原幂等键至多补偿一次，未决保持 UNKNOWN；200 任务争 50 库存时 0 超卖。"
    ),
}


def _replace(paragraph, label: str, text: str) -> None:
    if len(paragraph.runs) < 2:
        raise RuntimeError(f"bullet has no body run: {label}")
    paragraph.runs[0].text = label
    paragraph.runs[1].text = text
    for run in paragraph.runs[2:]:
        run.text = ""


def build() -> None:
    build_resume_v12.OUTPUT = OUTPUT
    build_resume_v12.build()

    doc = Document(OUTPUT)
    doc.core_properties.title = "袁明珠 - Agent 开发工程师（Java）实习简历 V13"
    doc.core_properties.subject = "岗位定向一页中文简历（最终证据版）"

    found = set()
    for paragraph in doc.paragraphs:
        if not paragraph.runs:
            continue
        label = paragraph.runs[0].text
        if label in BULLETS:
            _replace(paragraph, label, BULLETS[label])
            found.add(label)

    if found != set(BULLETS):
        raise RuntimeError(f"V13 replacement incomplete: missing={sorted(set(BULLETS) - found)}")

    project_labels = {
        "Agent 运行时：",
        "持久化执行与故障恢复：",
        "上下文与指代：",
        "检索与数据评测：",
        "商家评论 RAG：",
        "评测工程化：",
        "Multi-Agent：",
        "长期记忆：",
        "交易闭环：",
        "库存与订单补偿：",
        "异步投影：",
        "缓存与限流：",
        "架构取舍：",
    }
    for paragraph in doc.paragraphs:
        if paragraph.runs and paragraph.runs[0].text in project_labels:
            paragraph.paragraph_format.space_after = Pt(0.5)
            paragraph.paragraph_format.line_spacing = 0.98
            for run in paragraph.runs:
                run.font.size = Pt(9.0)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
