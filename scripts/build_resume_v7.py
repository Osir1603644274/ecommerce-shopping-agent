"""Build the evidence-bound, black-and-white one-page Chinese resume V7."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUTPUT = Path(r"C:\Users\ming\Desktop\袁明珠_简历\袁明珠_Agent开发工程师Java_岗位定向简历_V7.docx")
BODY_CN = "宋体"
HEAD_CN = "黑体"
LATIN = "Times New Roman"


def set_run_font(run, size: float, *, bold: bool = False, east_asia: str = BODY_CN, latin: str = LATIN) -> None:
    run.bold = bold
    run.font.name = latin
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)
    color = run._element.get_or_add_rPr().find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        run._element.get_or_add_rPr().append(color)
    color.set(qn("w:val"), "000000")


def set_keep(paragraph, *, next_paragraph: bool = False) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    keep_lines = OxmlElement("w:keepLines")
    ppr.append(keep_lines)
    if next_paragraph:
        keep_next = OxmlElement("w:keepNext")
        ppr.append(keep_next)


def add_bottom_border(paragraph, size: int = 8, space: int = 2) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)


def add_hanging_bullet_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    existing_abstract = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(existing_abstract, default=0) + 1
    num_id = max(existing_num, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "280")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "280")
    indent.set(qn("w:hanging"), "180")
    ppr.append(tabs)
    ppr.append(indent)
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    rpr.append(rfonts)
    for node in (start, num_fmt, lvl_text, lvl_jc, ppr, rpr):
        level.append(node)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def apply_bullet(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)
    ppr.append(num_pr)


def add_section(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4.8)
    p.paragraph_format.space_after = Pt(3.2)
    p.paragraph_format.line_spacing = 1.0
    set_keep(p, next_paragraph=True)
    add_bottom_border(p, size=8, space=1)
    set_run_font(p.add_run(text), 11.2, bold=True, east_asia=HEAD_CN, latin="Arial")


def add_subhead(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3.4)
    p.paragraph_format.space_after = Pt(1.6)
    p.paragraph_format.line_spacing = 1.0
    set_keep(p, next_paragraph=True)
    set_run_font(p.add_run(text), 9.9, bold=True, east_asia=HEAD_CN, latin="Arial")


def add_left_right(doc: Document, left: str, right: str, *, size: float = 10.0, bold_left: bool = True, bold_right: bool = False, before: float = 0, after: float = 1.0) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.tab_stops.add_tab_stop(Cm(18.2), WD_TAB_ALIGNMENT.RIGHT)
    set_keep(p)
    set_run_font(p.add_run(left), size, bold=bold_left, east_asia=HEAD_CN if bold_left else BODY_CN)
    p.add_run("\t")
    set_run_font(p.add_run(right), size, bold=bold_right)


def add_bullet(doc: Document, num_id: int, label: str, text: str, *, size: float = 10.4) -> None:
    p = doc.add_paragraph()
    apply_bullet(p, num_id)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.3)
    p.paragraph_format.line_spacing = 1.15
    set_keep(p)
    set_run_font(p.add_run(label), size, bold=True, east_asia=HEAD_CN)
    set_run_font(p.add_run(text), size)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(0.9)
    section.left_margin = Cm(1.18)
    section.right_margin = Cm(1.18)
    section.header_distance = Cm(0.3)
    section.footer_distance = Cm(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(10.4)
    normal._element.rPr.rFonts.set(qn("w:ascii"), LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CN)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15

    doc.core_properties.title = "袁明珠 - Agent 开发工程师（Java）实习简历"
    doc.core_properties.subject = "岗位定向一页中文简历"
    doc.core_properties.author = "袁明珠"
    doc.core_properties.keywords = "Java, Agent, ReAct, LangGraph, Spring Boot, Redis, MySQL"

    bullet_num = add_hanging_bullet_numbering(doc)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(1.2)
    name.paragraph_format.line_spacing = 1.0
    set_run_font(name.add_run("袁明珠"), 20.0, bold=True, east_asia=HEAD_CN, latin="Arial")

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(1.6)
    contact.paragraph_format.line_spacing = 1.0
    set_run_font(contact.add_run("Agent 开发工程师（Java）实习  |  北京  |  17559557087  |  17559557087@163.com"), 10.0, east_asia=BODY_CN, latin="Arial")

    add_section(doc, "教育经历")
    add_left_right(doc, "北京航空航天大学  |  电子信息  硕士", "2025.09 - 2028.06", size=10.6)
    add_left_right(doc, "河北工业大学  |  电子信息工程  本科", "2021.09 - 2025.06", size=10.6)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.line_spacing = 1.0
    set_run_font(p.add_run("本科专业排名 1/126，推免至北京航空航天大学"), 10.0)

    add_section(doc, "项目经历")
    add_left_right(doc, "电商购物 Agent 与 Java 交易平台  |  个人项目", "2026.07 - 至今", size=10.3, after=0.8)
    stack = doc.add_paragraph()
    stack.paragraph_format.space_after = Pt(0.8)
    stack.paragraph_format.line_spacing = 1.0
    set_run_font(stack.add_run("技术栈："), 9.7, bold=True, east_asia=HEAD_CN)
    set_run_font(stack.add_run("Python / FastAPI / LangGraph / MCP / Java 17 / Spring Boot / Spring Cloud / MySQL / Redis / Kafka / Elasticsearch"), 9.7, latin="Arial")
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(0.8)
    intro.paragraph_format.line_spacing = 1.02
    set_run_font(intro.add_run("面向多轮商品检索、比较与交易确认：Agent 负责低风险搜推决策，Java 后端持有商品、库存与订单权威写边界。"), 10.2)

    add_subhead(doc, "Agent 运行时、Context 与检索")
    add_bullet(doc, bullet_num, "受约束 ReAct：", "将固定 PAE 演进为“模型选择签名动作＋服务端执行/校验”，保留 fixed_v1 回滚；24 场景安全门 24/24、工具身份失败 0。")
    add_bullet(doc, bullet_num, "Context / 指代：", "实现 typed ContextCompiler 与服务端 ReferenceContext 收据；25 组冻结长对话两臂 25/25 等价，Prompt Token 降低26.43%、总 Token 降低22.92%；完成 4 轮真实网页链路，相关回归 577/577。")
    add_bullet(doc, bullet_num, "Checkpoint：", "以 Redis Checkpoint＋ToolInbox logical slot／lease／fence 恢复工具链；130/130 恢复、12/12 篡改拒绝，重复/缺失副作用与旧 Worker 覆盖均为 0，进程内 RTO P50/P95 42.182/75.475 ms。")
    add_bullet(doc, bullet_num, "检索选型：", "在 439 商品、24 Query 上比较 ES、Dense、RRF、Cross-Encoder 与 LLM 重排；混合方案 nDCG@10 +1.45pp 但 Recall@50 -5.01pp，故保留 ES Standard Top50＋MySQL 核验＋确定性重排，硬约束违规 0/831。")
    add_bullet(doc, bullet_num, "Multi-Agent：", "落地 Shopping Coordinator→只读 EvidenceResearchAgent→确定性合并，以 task/revision/CandidateScope 哈希与 Redis 原子 claim 约束通信；公开开发集 research 5/9→9/9，冻结合成留出 8/8、P95 5.50s、子观测泄漏 0。")
    add_bullet(doc, bullet_num, "记忆 / Skill / MCP：", "构建 MySQL 权威＋Redis 投影的可治理长期记忆，真实三会话链路 42/42；实现 10 路 UNJUDGED 候选池、显式 Skill 与本地 STDIO MCP，共用确定性核心并保持字节一致。")

    add_subhead(doc, "Java 交易后端与 Spring Cloud")
    add_bullet(doc, bullet_num, "权限与交易：", "Agent 不持有订单写能力；Java 回查 JWT 与明确确认，以 GETDEL＋Idempotency-Key 防重复消费/建单，Redis Lua＋数据库唯一边界保证一人一单。")
    add_bullet(doc, bullet_num, "异步一致性：", "订单与 Outbox 同事务，Kafka→Inbox→Elasticsearch 以 projection cursor／external_gte 处理重复、乱序与旧版本；异常路径显式回到 MySQL 权威。")
    add_bullet(doc, bullet_num, "缓存与并发：", "实现 Caffeine＋Redis 两级缓存及跨实例失效；200 个客户端任务争 50 库存得到 50 成功、150 售罄、0 超卖，Java 后端回归 197/197。")
    add_bullet(doc, bullet_num, "有界微服务拆分：", "拆分 Gateway、Catalog/Search、Trade，引入 Feign、Bulkhead(8) 与 CircuitBreaker；正式场景 10/10。对照发现拆分 P50 16.532 ms（单体 8.078 ms）、启动 45.068 s（单体 17.425 s），故保留模块化单体默认。")

    add_section(doc, "专业技能")
    add_bullet(doc, bullet_num, "Java / Spring：", "熟悉 Java 集合、异常、并发与 JVM 基础；熟练使用 Spring Boot、IOC/AOP、Spring Security、MyBatis，具备 Spring Cloud 项目实践。", size=10.2)
    add_bullet(doc, bullet_num, "数据与中间件：", "熟悉 MySQL 事务、隔离级别、索引与锁；掌握 Redis 缓存/Lua/ZSET/Stream，具备 Kafka、Elasticsearch、Caffeine 实践。", size=10.2)
    add_bullet(doc, bullet_num, "Agent 工程：", "熟悉 ReAct/LangGraph、Tool Calling、Checkpoint、Context、RAG、Multi-Agent、长期记忆与 MCP。", size=10.2)
    add_bullet(doc, bullet_num, "工程工具：", "使用 Maven、Docker Compose、JUnit、pytest 完成故障注入与证据化评测；Agent 全量回归 3593 passed、12 skipped、0 failed。", size=10.2)

    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(str(OUTPUT))


if __name__ == "__main__":
    build()
