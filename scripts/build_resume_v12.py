"""Build the black-and-white one-page Chinese resume V12.

V12 keeps the complete Java backend section from V11 and strengthens the
Agent-side evidence hierarchy: broad/deep/external retrieval datasets,
durable execution terminology, and bounded Skill/MCP engineering claims.
"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.text.paragraph import Paragraph

import build_resume_v11


OUTPUT = Path(
    r"C:\Users\ming\Desktop\袁明珠_简历\袁明珠_Agent开发工程师Java_岗位定向简历_V12.docx"
)


INTRO = (
    "面向多品类电商检索与二手手机深度导购，打通搜索、比较、受控下单与订单查询；"
    "Python Agent 负责决策，Java 后端持有身份、库存、订单与支付权威。"
)


BULLETS = {
    "Agent 运行时：": (
        "针对固定 PAE 难以适配灵活追问的问题，先以 LangGraph 状态图、Interrupt 与 Checkpoint 建立可暂停/恢复底座，"
        "再将决策层收敛为受约束 ReAct；24 场景每臂 33 轮，差异回答双评审盲评 ReAct/fixed/tie=7/0/3，并保留 fixed_v1 回滚。"
    ),
    "持久化执行与故障恢复：": (
        "区分澄清挂起、用户主动暂停与进程崩溃续跑，以 Redis Checkpoint、ToolInbox、租约及 fencing token 防重复执行和旧 Worker 覆盖；"
        "130/130 恢复、12/12 篡改拒绝，进程内 RTO P95 约 75ms；跨系统 Exactly-once 不作声明。"
    ),
    "上下文与指代：": (
        "针对全量历史随轮次膨胀和商品错指，构建类型化 ContextCompiler 与 ReferenceContext 服务端收据；"
        "25 组冻结合成长历史保持 25/25 结果等价，总 Token 下降 22.9%，跨会话、旧任务及越界引用失败关闭。"
    ),
    "检索与数据评测：": (
        "以 KuaiSearch 46,079 文档、507 条真实 Query、12 品类建立广度基准，并用 439 件二手手机、24 条复杂 Query 深测约束检索；"
        "对比 ES/Dense/RRF/模型重排后保留 ES Standard Top50，Amazon ESCI 50 Query、1,036 标注 pair 上重排 nDCG@10=0.7069（随机 0.5283）。"
    ),
    "商家评论 RAG：": (
        "针对短评证据分散，采用结构化过滤、双路召回与内容重排，将离线 nDCG@5 从 0.368 提升至 0.788；"
        "因延迟与忠实度门未通过，仅保留 Shadow，不进入默认链路。"
    ),
    "Multi-Agent：": (
        "仅在证据缺口时触发只读 EvidenceResearchAgent，消息绑定任务版本与 CandidateScope，Redis 原子合并后由父 Agent 确定性渲染；"
        "开发集 5/9→9/9、冻结合成留出 8/8、原始子观测泄漏 0，异常回退单 Agent。"
    ),
    "长期记忆：": (
        "为防点击和普通对话误写偏好，设计“异步候选→用户确认→MySQL 权威版本链→Redis 投影”，"
        "支持跨 Session 召回、纠正、停用与删除；真实 MySQL/Redis/ES/Java/模型三 Session 链路完成 42/42 有界验证。"
    ),
}


EVALUATION = (
    "将 10 路检索候选池构建封装为显式 Skill 与本地 STDIO MCP，共用确定性内核及 manifest/receipt/SHA 证据链；"
    "133 项 SHA 复核零失败，显式 Skill/MCP 输出字节一致；隐式加载与质量提升仍为 HOLD。"
)


def _replace_bullet(paragraph, label: str, text: str) -> None:
    if len(paragraph.runs) < 2:
        raise RuntimeError(f"bullet has no body run: {label}")
    paragraph.runs[0].text = label
    paragraph.runs[1].text = text
    for run in paragraph.runs[2:]:
        run.text = ""


def _insert_before(paragraph, label: str, text: str):
    clone = deepcopy(paragraph._p)
    paragraph._p.addprevious(clone)
    inserted = Paragraph(clone, paragraph._parent)
    _replace_bullet(inserted, label, text)
    return inserted


def build() -> None:
    build_resume_v11.OUTPUT = OUTPUT
    build_resume_v11.build()

    doc = Document(OUTPUT)
    doc.core_properties.title = "袁明珠 - Agent 开发工程师（Java）实习简历 V12"
    doc.core_properties.subject = "岗位定向一页中文简历（数据与评测增强版）"

    found_intro = False
    found = set()
    inserted_evaluation = False

    for paragraph in list(doc.paragraphs):
        if paragraph.text.startswith("面向二手 3C 多轮导购"):
            paragraph.runs[0].text = INTRO
            for run in paragraph.runs[1:]:
                run.text = ""
            found_intro = True
            continue
        if not paragraph.runs:
            continue
        label = paragraph.runs[0].text
        if label == "可靠执行：":
            _replace_bullet(
                paragraph,
                "持久化执行与故障恢复：",
                BULLETS["持久化执行与故障恢复："],
            )
            found.add("持久化执行与故障恢复：")
        elif label == "检索 / RAG 选型：":
            _replace_bullet(paragraph, "检索与数据评测：", BULLETS["检索与数据评测："])
            found.add("检索与数据评测：")
        elif label == "Multi-Agent：":
            _insert_before(paragraph, "商家评论 RAG：", BULLETS["商家评论 RAG："])
            found.add("商家评论 RAG：")
            _insert_before(paragraph, "评测工程化：", EVALUATION)
            inserted_evaluation = True
            _replace_bullet(paragraph, label, BULLETS[label])
            found.add(label)
        elif label in BULLETS:
            _replace_bullet(paragraph, label, BULLETS[label])
            found.add(label)

    # Keep one page despite the extra evidence bullet while preserving readable
    # black-and-white typography and the full five-bullet Java backend section.
    project_labels = set(BULLETS) | {
        "评测工程化：",
        "交易闭环：",
        "库存与订单补偿：",
        "异步投影：",
        "缓存与限流：",
        "架构取舍：",
    }
    for paragraph in doc.paragraphs:
        if paragraph.runs and paragraph.runs[0].text in project_labels:
            paragraph.paragraph_format.space_after = Pt(0.7)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.size = Pt(9.25)

    expected = set(BULLETS)
    if not found_intro or found != expected or not inserted_evaluation:
        raise RuntimeError(
            "V12 replacement incomplete: "
            f"intro={found_intro}, missing={sorted(expected - found)}, "
            f"evaluation={inserted_evaluation}"
        )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
