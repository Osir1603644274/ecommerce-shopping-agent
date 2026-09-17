from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = Path("outputs/resume/袁明珠-RAG检索算法实习-阿里千问.docx")

BODY_CN = "SimSun"
HEAD_CN = "SimHei"
LATIN = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(65, 65, 65)


def set_run_font(run, size=9.5, bold=False, color=BLACK, east_asia=BODY_CN):
    run.font.name = LATIN
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), LATIN)
    fonts.set(qn("w:hAnsi"), LATIN)
    fonts.set(qn("w:eastAsia"), east_asia)


def set_style_font(style, size=9.5, bold=False, east_asia=BODY_CN):
    style.font.name = LATIN
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = BLACK
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), LATIN)
    fonts.set(qn("w:hAnsi"), LATIN)
    fonts.set(qn("w:eastAsia"), east_asia)


def add_paragraph_style(doc, name, size, bold=False, before=0, after=0, line=1.0, east_asia=BODY_CN):
    style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    set_style_font(style, size=size, bold=bold, east_asia=east_asia)
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    return style


def set_keep(paragraph, keep_next=False):
    ppr = paragraph._p.get_or_add_pPr()
    keep_lines = ppr.find(qn("w:keepLines"))
    if keep_lines is None:
        ppr.append(OxmlElement("w:keepLines"))
    if keep_next and ppr.find(qn("w:keepNext")) is None:
        ppr.append(OxmlElement("w:keepNext"))


def set_bottom_border(paragraph, size=8, space=3):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), "000000")
    pbdr.append(bottom)


def create_bullet_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)

    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "430")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "430")
    ind.set(qn("w:hanging"), "260")
    ppr.append(ind)
    lvl.append(ppr)

    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Symbol")
    rfonts.set(qn("w:hAnsi"), "Symbol")
    rpr.append(rfonts)
    lvl.append(rpr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def apply_bullet(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    ppr.append(num_pr)


def add_section_heading(doc, text):
    p = doc.add_paragraph(style="ResumeSection")
    set_run_font(p.add_run(text), size=11.5, bold=True, east_asia=HEAD_CN)
    set_bottom_border(p, size=8, space=3)
    set_keep(p, keep_next=True)
    return p


def add_timed_line(doc, left_bold, left_rest, date_text, style="ResumeEntry"):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
    set_run_font(p.add_run(left_bold), size=9.9, bold=True, east_asia=HEAD_CN)
    set_run_font(p.add_run(left_rest), size=9.5, color=GRAY)
    set_run_font(p.add_run("\t" + date_text), size=9.3, color=GRAY)
    set_keep(p)
    return p


def add_labeled_line(doc, label, content, style="ResumeLine"):
    p = doc.add_paragraph(style=style)
    set_run_font(p.add_run(label), size=9.5, bold=True, east_asia=HEAD_CN)
    set_run_font(p.add_run(content), size=9.5)
    set_keep(p)
    return p


def add_bullet(doc, num_id, label, content):
    p = doc.add_paragraph(style="ResumeBullet")
    apply_bullet(p, num_id)
    set_run_font(p.add_run(label), size=9.45, bold=True, east_asia=HEAD_CN)
    set_run_font(p.add_run(content), size=9.45)
    set_keep(p)
    return p


def configure_doc():
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    section.header_distance = Cm(0.4)
    section.footer_distance = Cm(0.4)

    normal = doc.styles["Normal"]
    set_style_font(normal, size=9.5)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.10

    add_paragraph_style(doc, "ResumeName", 19.0, bold=True, after=1.5, east_asia=HEAD_CN)
    add_paragraph_style(doc, "ResumeHeader", 9.3, after=1.2, line=1.0)
    add_paragraph_style(doc, "ResumeSection", 11.5, bold=True, before=5.0, after=3.0, line=1.0, east_asia=HEAD_CN)
    add_paragraph_style(doc, "ResumeEntry", 9.5, after=1.4, line=1.0)
    add_paragraph_style(doc, "ResumeLine", 9.5, after=1.5, line=1.12)
    add_paragraph_style(doc, "ResumeBullet", 9.45, after=1.8, line=1.16)
    return doc


def build():
    doc = configure_doc()
    bullet_num_id = create_bullet_numbering(doc)

    p = doc.add_paragraph(style="ResumeName")
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
    set_run_font(p.add_run("袁明珠"), size=19.0, bold=True, east_asia=HEAD_CN)
    set_run_font(p.add_run("\t17559557087  |  1603644274@qq.com  |  北京"), size=9.3, color=GRAY)

    p = doc.add_paragraph(style="ResumeHeader")
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
    set_run_font(p.add_run("求职意向：RAG 检索算法实习生"), size=10.0, bold=True, east_asia=HEAD_CN)
    set_run_font(p.add_run("\t2028 届  |  每周至少 4 天  |  可实习 6 个月"), size=9.2, color=GRAY)

    add_section_heading(doc, "教育经历")
    add_timed_line(doc, "北京航空航天大学", "  |  电子信息  |  硕士在读", "2025.09—2028.06")
    add_timed_line(doc, "河北工业大学", "  |  电子信息工程  |  本科", "2021.09—2025.06")

    add_section_heading(doc, "项目经历")
    add_timed_line(
        doc,
        "二手 3C Agentic Product Search",
        "  |  RAG 检索与自动化评测系统  |  个人项目",
        "2026.06—至今",
    )
    add_labeled_line(
        doc,
        "技术栈：",
        "Python + FastAPI + LangGraph + PyTorch/Transformers + Elasticsearch + Qdrant + Redis + Java/Spring Boot + MySQL + Kafka",
    )
    add_labeled_line(
        doc,
        "项目简介：",
        "面向二手 3C 导购的对话式 Agent，在受控商品域内完成自然语言需求理解、候选检索与排序、权威事实核验、多轮决策和自动化评测。",
    )

    add_bullet(
        doc,
        bullet_num_id,
        "Agent 与 Query 理解：",
        "将自然语言需求沉淀为带 revision 的 TaskState，使用 Planner、Executor、Validator、Replanner 形成检索、详情核验、比较、澄清和恢复闭环，避免模型直接生成未经验证的商品事实。",
    )
    add_bullet(
        doc,
        bullet_num_id,
        "混合检索与重排：",
        "搭建并对比 Elasticsearch/BM25、BGE Dense Retrieval、RRF 融合与 Cross-Encoder 重排；分离品牌、型号、内存、成色等硬约束和相关性排序，并回查 Java/MySQL 权威事实。",
    )
    add_bullet(
        doc,
        bullet_num_id,
        "RAG 与效果分析：",
        "基于 Qdrant、BGE Embedding 和 BM25 实现带来源引用的评论 RAG，将检索命中与回答忠实性拆分评估；真实 Yelp 评论 25 题 Hit@3 为 60.0%，据失败样本确定 Hybrid Search、Query Rewrite、metadata filter 与 Rerank 优化方向。",
    )
    add_bullet(
        doc,
        bullet_num_id,
        "自动化数据与评测闭环：",
        "构建候选池、标注/qrels、冻结测试及 Benchmark/Harness，统一输出 Hit@K、Recall、MRR、nDCG、硬约束违规和端到端延迟；覆盖多轮澄清、无结果、约束冲突、依赖故障与状态恢复，并生成可追溯 trace 和评分产物。",
    )

    add_timed_line(
        doc,
        "Agentic Commerce Backend",
        "  |  Java 搜索与交易后端",
        "同项目子系统",
    )
    add_labeled_line(
        doc,
        "技术栈：",
        "Java + Spring Boot + MyBatis + MySQL + Redis + Kafka + Elasticsearch + Docker",
    )
    add_bullet(
        doc,
        bullet_num_id,
        "搜索事实边界：",
        "使用 Elasticsearch 多字段检索与硬条件 filter，设置超时熔断和 MySQL fallback；ES 仅返回排序 ID，Service 层重新核验价格、状态和约束，避免陈旧索引进入最终决策。",
    )
    add_bullet(
        doc,
        bullet_num_id,
        "缓存与一致性：",
        "构建 Caffeine L1 + Redis L2 cache-aside，处理空值穿透、热点击穿、TTL 抖动和损坏值回退，并通过事件驱动失效实现跨实例 L1 有界一致性。",
    )
    add_bullet(
        doc,
        bullet_num_id,
        "可靠交易链路：",
        "使用订单/支付状态机与 Idempotency-Key 防止重复写入，通过 Transactional Outbox + Kafka + Inbox 实现至少一次投递、消费幂等和可审计失败恢复。",
    )

    add_section_heading(doc, "专业技能")
    add_bullet(doc, bullet_num_id, "编程与工程：", "使用 Python、Java、SQL，能够基于 FastAPI、Spring Boot、Pydantic 与 Docker 完成服务开发和接口联调。")
    add_bullet(doc, bullet_num_id, "大模型与 Agent：", "使用 LangGraph、Function Calling、Prompt/Tool Schema、ReAct/PAE、多轮 TaskState 和 Agent Harness。")
    add_bullet(doc, bullet_num_id, "检索与评测：", "使用 Elasticsearch/BM25、Qdrant、Dense Retrieval、RRF、Cross-Encoder；掌握 Hit@K、Recall@K、MRR、nDCG 和延迟评测。")
    add_bullet(doc, bullet_num_id, "模型与数据：", "使用 PyTorch/Transformers 进行 Embedding 与 Reranker 推理，使用 MySQL、Redis、Kafka 构建数据与状态链路。")

    doc.core_properties.title = "袁明珠 - RAG 检索算法实习简历"
    doc.core_properties.subject = "阿里千问日常算法实习生投递版"
    doc.core_properties.author = "袁明珠"
    doc.core_properties.keywords = "RAG, 检索算法, Agent, Elasticsearch, BM25, Qdrant"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
