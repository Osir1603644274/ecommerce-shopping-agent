"""Build the balanced black-and-white one-page Chinese resume V10."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Cm, Pt

from build_resume_v9 import (
    BODY_CN,
    HEAD_CN,
    LATIN,
    add_bottom_border,
    add_bullet_numbering,
    apply_bullet,
    set_keep,
    set_run_font,
)


OUTPUT = Path(r"C:\Users\ming\Desktop\袁明珠_简历\袁明珠_Agent开发工程师Java_岗位定向简历_V10.docx")


def add_section(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4.0)
    p.paragraph_format.space_after = Pt(2.8)
    p.paragraph_format.line_spacing = 1.0
    set_keep(p, next_paragraph=True)
    add_bottom_border(p, size=8, space=1)
    set_run_font(p.add_run(text), 11.2, bold=True, east_asia=HEAD_CN, latin="Arial")


def add_subhead(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3.0)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.0
    set_keep(p, next_paragraph=True)
    set_run_font(p.add_run(text), 10.2, bold=True, east_asia=HEAD_CN, latin="Arial")


def add_left_right(doc: Document, left: str, right: str, *, size: float = 10.1, after: float = 0.8) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
    set_keep(p)
    set_run_font(p.add_run(left), size, bold=True, east_asia=HEAD_CN)
    p.add_run("\t")
    set_run_font(p.add_run(right), size)


def add_label_line(doc: Document, label: str, text: str, *, size: float = 9.6, after: float = 0.8) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    set_keep(p)
    set_run_font(p.add_run(label), size, bold=True, east_asia=HEAD_CN, latin="Arial")
    set_run_font(p.add_run(text), size, latin="Arial")


def add_bullet(doc: Document, num_id: int, label: str, text: str, *, size: float = 9.7, after: float = 1.25) -> None:
    p = doc.add_paragraph()
    apply_bullet(p, num_id)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.widow_control = True
    set_keep(p)
    set_run_font(p.add_run(label), size, bold=True, east_asia=HEAD_CN)
    set_run_font(p.add_run(text), size)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    # V10 named override: A4 resume geometry. Keep a readable 9.7 pt body and
    # use restrained 1.45 cm side margins instead of compressing the type.
    section.top_margin = Cm(1.05)
    section.bottom_margin = Cm(0.95)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    section.header_distance = Cm(0.25)
    section.footer_distance = Cm(0.25)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(9.7)
    normal._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii", LATIN)
    normal._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi", LATIN)
    normal._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", BODY_CN)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.05

    doc.core_properties.title = "袁明珠 - Agent 开发工程师（Java）实习简历 V10"
    doc.core_properties.subject = "岗位定向一页中文简历（平衡版）"
    doc.core_properties.author = "袁明珠"
    doc.core_properties.keywords = "Java, Agent, ReAct, LangGraph, Spring Boot, Redis, MySQL, RAG, Multi-Agent"

    bullet_num = add_bullet_numbering(doc)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(0.8)
    name.paragraph_format.line_spacing = 1.0
    set_run_font(name.add_run("袁明珠"), 19.5, bold=True, east_asia=HEAD_CN, latin="Arial")

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(1.0)
    contact.paragraph_format.line_spacing = 1.0
    set_run_font(
        contact.add_run("Agent 开发工程师（Java）实习生  |  北京  |  17559557087  |  17559557087@163.com"),
        9.65,
        east_asia=BODY_CN,
        latin="Arial",
    )

    add_section(doc, "教育经历")
    add_left_right(doc, "北京航空航天大学  |  电子信息  硕士", "2025.09 - 2028.06", size=10.35)
    add_left_right(doc, "河北工业大学  |  电子信息工程  本科", "2021.09 - 2025.06", size=10.35)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0.6)
    p.paragraph_format.line_spacing = 1.0
    set_run_font(p.add_run("本科专业排名 1/126，推免至北京航空航天大学"), 9.8)

    add_section(doc, "项目经历")
    add_left_right(doc, "电商购物 Agent 与 Java 交易平台  |  个人项目", "2026.07 - 至今", size=10.4, after=0.8)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(0.8)
    intro.paragraph_format.line_spacing = 1.02
    set_run_font(
        intro.add_run("面向多轮商品检索、证据核验与受控交易：Python Agent 负责搜推决策，Java 后端持有身份、库存、订单与支付写边界。"),
        9.75,
    )

    add_subhead(doc, "智能购物 Agent")
    add_label_line(doc, "技术栈：", "Python / FastAPI / LangGraph / Redis / MySQL / Elasticsearch / MCP")
    add_bullet(
        doc,
        bullet_num,
        "Agent 运行时：",
        "针对固定 Planner-Executor-Validator 流程推进僵硬的问题，将其收敛为受约束 ReAct；模型仅选择服务端签发动作，执行与校验依赖服务端任务状态及候选范围。24/24 安全场景通过，5 个差异场景双人盲评共 10 票，ReAct/fixed/tie=7/0/3，并保留 fixed_v1 回滚。",
    )
    add_bullet(
        doc,
        bullet_num,
        "可靠执行：",
        "使用 Redis Checkpoint、幂等工具收件箱（ToolInbox）及租约/fencing token，处理工具成功后崩溃造成的重复执行和旧 Worker 覆盖；130/130 恢复、12/12 篡改拒绝，重复或缺失副作用为 0，进程内 RTO P95 约 75ms。",
    )
    add_bullet(
        doc,
        bullet_num,
        "上下文与指代：",
        "针对全量历史随轮次增长及商品错指风险，构建类型化 ContextCompiler 与服务端引用收据，绑定会话、任务版本、候选范围和真实展示顺序；25 组冻结合成长历史中两臂均 25/25 等价，总 Token 下降 22.9%。",
    )
    add_bullet(
        doc,
        bullet_num,
        "检索 / RAG 选型：",
        "在 439 商品、24 Query 上对比 ES、Dense、RRF 及模型重排；混合方案 nDCG@10 +1.45pp 但 Recall@50 -5.01pp，故采用 ES Standard Top50＋MySQL 核验＋确定性重排。早期商家评论 RAG 以双路召回和内容重排将 nDCG@5 从 0.368 提至 0.788，但因延迟与忠实度未过门，仅保留 Shadow。",
    )
    add_bullet(
        doc,
        bullet_num,
        "Multi-Agent：",
        "仅在证据缺口时触发只读 EvidenceResearchAgent，消息绑定任务版本与候选范围，经 Redis 原子合并后由父 Agent 确定性渲染；公开开发修复 5/9→9/9、冻结合成留出 8/8，原始子观测泄漏 0，异常回退单 Agent。",
    )
    add_bullet(
        doc,
        bullet_num,
        "长期记忆：",
        "为避免把点击和普通对话误写为长期偏好，设计“异步候选→用户确认→MySQL 权威版本链→Redis 投影”流程，支持跨 Session 召回、纠正、停用与删除；真实 MySQL/Redis/ES/Java/模型三 Session 链路完成 42/42 有界验证。",
    )

    add_subhead(doc, "Java 电商交易后端")
    add_label_line(
        doc,
        "技术栈：",
        "Java 17 / Spring Boot / Spring Security / MyBatis / MySQL / Redis / Kafka / Elasticsearch / Caffeine / Spring Cloud",
    )
    add_bullet(
        doc,
        bullet_num,
        "交易写边界：",
        "按商品、库存、订单、支付/退款与搜索划分模块化单体；Agent 不持有订单写工具，下单由 Java 回查 JWT、校验候选范围并要求用户二次确认，一次性确认令牌与 Idempotency-Key 分别防止重复确认和重复建单。",
    )
    add_bullet(
        doc,
        bullet_num,
        "库存与订单补偿：",
        "Redis Lua 原子预扣/一人一单，MySQL 条件扣减与唯一索引提供最终边界；200 个客户端任务争 50 库存得到 50 成功、150 售罄、0 超卖。订单过期迁移、库存/优惠券释放及 Outbox 同事务，2 个并发 Worker 仅产生 1 次迁移和 1 条事件。",
    )
    add_bullet(
        doc,
        bullet_num,
        "异步投影：",
        "订单/商品变更与 Outbox 同事务，经 Kafka→Inbox→Elasticsearch 更新搜索投影，并以版本游标和 external_gte 处理重复、乱序与旧版本事件，异常路径回查 MySQL 权威数据。",
    )
    add_bullet(
        doc,
        bullet_num,
        "缓存与限流：",
        "商品详情采用 Caffeine＋Redis 两级缓存，事务提交后删除 Redis 并广播跨实例 L1 失效；Redis ZSET＋Lua 按 IP、JWT 用户和归一化接口实现滑动窗口限流，100 并发、额度 20 时准确放行 20、拒绝 80。",
    )
    add_bullet(
        doc,
        bullet_num,
        "架构取舍：",
        "验证 Gateway、Catalog/Search、Trade 拆分形态，引入 Feign、Bulkhead 与 CircuitBreaker；同源对照中拆分 P50 由 8.08ms 增至 16.53ms、启动由 17.43s 增至 45.07s，因此默认保留模块化单体。",
    )

    add_section(doc, "专业技能")
    add_bullet(doc, bullet_num, "Java / Spring：", "Java 17、集合与并发、JVM 基础、Spring Boot、Spring Security、MyBatis、事务、IOC/AOP、Spring Cloud。", size=9.7, after=0.7)
    add_bullet(doc, bullet_num, "数据与中间件：", "MySQL 事务/索引/锁，Redis 缓存、Lua、ZSET、Stream，Kafka、Elasticsearch、Caffeine。", size=9.7, after=0.7)
    add_bullet(doc, bullet_num, "Agent 工程：", "LangGraph、ReAct、Tool Calling、RAG、Multi-Agent、Checkpoint、Context/Memory、Skill/MCP；Maven、Docker Compose、JUnit、pytest。", size=9.7, after=0)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
